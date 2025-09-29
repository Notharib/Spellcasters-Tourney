import os
import re
import sys
from docx import Document

def extract_comment_blocks(file_content):
    # Match triple-quoted blocks
    block_pattern = re.compile(r"(?:'''|\"\"\")(.*?)(?:'''|\"\"\")", re.DOTALL)
    blocks = block_pattern.findall(file_content)

    # Match consecutive hash-prefixed lines
    hash_lines = []
    current_block = []
    for line in file_content.splitlines():
        stripped = line.strip()
        if stripped.startswith("#"):
            content = stripped[1:].strip()
            current_block.append(content)
        else:
            if current_block:
                hash_lines.append("\n".join(current_block))
                current_block = []
    if current_block:
        hash_lines.append("\n".join(current_block))

    return blocks + hash_lines



def parse_comment_block(block):
    fields = {'Name': '', 'Parameters': '', 'Returns': '', 'Purpose': ''}
    for line in block.strip().split('\n'):
        line = line.strip()
        for key in fields:
            if line.startswith(f"{key}:"):
                fields[key] = line[len(key)+1:].strip()
    return fields

def process_python_file(filepath):
    with open(filepath, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    entries = []
    current_comment = []
    in_block = False
    block_delimiter = None

    for i, line in enumerate(lines):
        stripped = line.strip()

        # Start of triple-quoted block
        if not in_block and (stripped.startswith("'''") or stripped.startswith('"""')):
            in_block = True
            block_delimiter = stripped[:3]
            current_comment = [stripped[3:]] if len(stripped) > 3 else []
            continue

        # Inside triple-quoted block
        if in_block:
            if stripped.endswith(block_delimiter):
                current_comment.append(stripped[:-3] if len(stripped) > 3 else '')
                in_block = False
                block_delimiter = None
            else:
                current_comment.append(stripped)
            continue

        # Hash-prefixed comment line
        if stripped.startswith("#"):
            current_comment.append(stripped[1:].strip())
            continue

        # Function definition
        if stripped.startswith("def ") and current_comment:
            comment_text = "\n".join(current_comment).strip()
            parsed = parse_comment_block(comment_text)
            entries.append(parsed)
            current_comment = []

    return entries


def write_to_word(tables_by_file, output_path):
    doc = Document()
    for filename, entries in tables_by_file.items():
        doc.add_heading(f"File: {filename}", level=2)
        if entries:
            table = doc.add_table(rows=1, cols=4)
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Name'
            hdr_cells[1].text = 'Parameters'
            hdr_cells[2].text = 'Returns'
            hdr_cells[3].text = 'Purpose'
            for entry in entries:
                row_cells = table.add_row().cells
                row_cells[0].text = entry['Name']
                row_cells[1].text = entry['Parameters']
                row_cells[2].text = entry['Returns']
                row_cells[3].text = entry['Purpose']
        else:
            doc.add_paragraph("No formatted comments found.")
        doc.add_paragraph()  # spacing between tables
    doc.save(output_path)

def main(folder_path, output_docx):
    tables_by_file = {}
    for root, _, files in os.walk(folder_path):
        for filename in files:
            if filename.endswith('.py'):
                filepath = os.path.join(root, filename)
                relative_path = os.path.relpath(filepath, folder_path)
                entries = process_python_file(filepath)
                if entries:
                    tables_by_file[relative_path] = entries
    write_to_word(tables_by_file, output_docx)

# Example usage
if __name__ == "__main__":
    folder_path = os.path.dirname(sys.executable if getattr(sys, 'frozen', False) else __file__)
    output_docx = "output.docx"
    main(folder_path, output_docx)
